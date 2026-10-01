# -*- coding: utf-8 -*-
"""md → docx 转换（python-docx，无需 pandoc）
用法: python tools/md2docx.py <input.md> <output.docx>
支持的 md 元素：# 标题(1-4级)、- 列表、1. 有序列表、> 引用、**粗体**、表格、--- 分隔线、普通段落
"""
import sys, re, os

def add_runs(paragraph, text):
    """处理 **粗体** 行内标记"""
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part:
            paragraph.add_run(part)

def md2docx(md_path, docx_path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    # 基础样式
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    from docx.oxml.ns import qn
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    lines = open(md_path, encoding='utf-8').read().split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 分隔线
        if stripped == '---':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run('· · ·')
            r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            i += 1
            continue

        # 表格
        if stripped.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i+1].strip()):
            header = [c.strip() for c in stripped.strip('|').split('|')]
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith('|'):
                rows.append([c.strip() for c in lines[j].strip().strip('|').split('|')])
                j += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = 'Light Grid Accent 1'
            for ci, cell in enumerate(header):
                table.rows[0].cells[ci].text = re.sub(r'\*\*', '', cell)
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    if ci < len(header):
                        table.rows[ri+1].cells[ci].text = re.sub(r'\*\*', '', cell)
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9)
            i = j
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)', stripped)
        if m:
            level = len(m.group(1))
            text = re.sub(r'\*\*', '', m.group(2))
            h = doc.add_heading(text, level=level)
            for r in h.runs:
                r.font.name = 'Microsoft YaHei'
                if r.element.rPr and r.element.rPr.rFonts:
                    r.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            i += 1
            continue

        # 引用
        if stripped.startswith('>'):
            # 收集连续引用行
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                ql = lines[i].strip()[1:].strip()
                if ql:
                    quote_lines.append(ql)
                i += 1
            for ql in quote_lines:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_after = Pt(2)
                add_runs(p, ql)
                for r in p.runs:
                    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                    r.font.size = Pt(10)
            continue

        # 无序列表
        if re.match(r'^[-*]\s+', stripped):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, re.sub(r'^[-*]\s+', '', stripped))
            i += 1
            continue

        # 有序列表
        m = re.match(r'^(\d+)[.、]\s+(.*)', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(2))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"✅ {docx_path} ({os.path.getsize(docx_path)} bytes)")

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("用法: python tools/md2docx.py <input.md> <output.docx>")
        sys.exit(1)
    md2docx(sys.argv[1], sys.argv[2])
